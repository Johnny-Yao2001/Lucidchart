"""Rebuild assignment CSVs without editing the supplied Excel files.
Usage: python prepare_data.py /path/to/Lucidchart
Requires openpyxl. All operational records are fictional test fixtures.
"""
from pathlib import Path
from collections import Counter, defaultdict
from decimal import Decimal
import csv
import hashlib
import json
import re
import sys
from urllib.parse import urlparse
from openpyxl import load_workbook

OUT = Path(__file__).resolve().parent
SOURCE = Path(sys.argv[1]) if len(sys.argv) > 1 else Path('/workspace/Lucidchart')
DATA = {}
SCHEMA = {}
ISSUES = []
LINEAGE = []
RAW = []

def define(name, columns):
    SCHEMA[name] = columns.split()
    DATA[name] = []

define('Supplier', 'SupplierID SupplierName Website Email')
define('ProductCategory', 'CategoryID CategoryName')
define('ProductGroup', 'ProductGroupID GroupName ProductStream')
define('Product', 'ProductID CategoryID ProductGroupID SourceProductID Description UnitOfMeasure UnitPrice Discontinued MinTempC MaxTempC GTIN')
define('ProductSupplier', 'ProductSupplierID ProductID SupplierID SupplierProductCode IsPreferred')
define('Warehouse', 'WarehouseID WarehouseName City Address')
define('Worker', 'WorkerID WarehouseID WorkerName EmploymentType')
define('StorageLocation', 'LocationID WarehouseID LocationCode ZoneType MinTempC MaxTempC IsActive')
define('Customer', 'CustomerID CustomerName CustomerType Description')
define('CustomerSite', 'CustomerSiteID CustomerID SiteName Address StateRegion Postcode Country Phone Email Website')
define('PurchaseOrder', 'PurchaseOrderID SupplierID WarehouseID OrderedAt ExpectedAt Status')
define('PurchaseOrderLine', 'PurchaseOrderLineID PurchaseOrderID ProductID OrderedQty UnitCost')
define('Receipt', 'ReceiptID PurchaseOrderID WorkerID ReceivedAt Status')
define('ReceiptLine', 'ReceiptLineID ReceiptID PurchaseOrderLineID ReceivedQty')
define('InventoryLot', 'LotID ProductID LotCode ExpiryDate')
define('ReceiptLineLot', 'ReceiptLineLotID ReceiptLineID LotID ReceivedQty')
define('StockBalance', 'StockBalanceID LotID LocationID OnHandQty LastUpdatedAt')
define('TemperatureLog', 'TemperatureLogID LocationID RecordedAt TemperatureC')
define('SalesOrder', 'SalesOrderID CustomerSiteID WarehouseID OrderedAt PromisedDeliveryAt Status')
define('SalesOrderLine', 'SalesOrderLineID SalesOrderID ProductID RequestedQty')
define('PickLine', 'PickLineID SalesOrderLineID WorkerID StockBalanceID PickedQty CorrectQty PickedAt VerifiedAt')
define('Shipment', 'ShipmentID SalesOrderID DeliveryMethod CarrierName DispatchedAt DeliveredAt DeliveryTempC')
define('ShipmentLine', 'ShipmentLineID ShipmentID SalesOrderLineID ShippedQty')
define('ShipmentLineLot', 'ShipmentLineLotID ShipmentLineID LotID ShippedQty')

def add(table, *values):
    assert len(values) == len(SCHEMA[table]), (table, values)
    DATA[table].append(dict(zip(SCHEMA[table], values)))
    return values[0]

def clean(value):
    if value is None:
        return None
    text = re.sub(r'\s+', ' ', str(value)).strip()
    return text or None

def issue(file, sheet, row, field, original, action):
    ISSUES.append([file, sheet, row, field, original, action])

def source_rows(filename):
    workbook = load_workbook(SOURCE / filename, read_only=True, data_only=True)
    for sheet in workbook:
        rows = sheet.iter_rows(values_only=True)
        headers = [clean(x) or '' for x in next(rows)]
        for number, values in enumerate(rows, 2):
            if not any(clean(x) for x in values):
                continue
            raw = {h: v for h, v in zip(headers, values) if h}
            RAW.append({'file': filename, 'sheet': sheet.title, 'row': number, 'values': raw})
            yield sheet.title, number, {h: clean(v) for h, v in raw.items()}
    workbook.close()

def website(value, file, sheet, row):
    if not value:
        return None
    candidate = value if '://' in value else 'https://' + value
    parsed = urlparse(candidate)
    if parsed.scheme in ('http', 'https') and parsed.hostname and re.fullmatch(r'[A-Za-z0-9.-]+\.[A-Za-z]{2,}', parsed.hostname) and ' ' not in value:
        if candidate != value:
            issue(file, sheet, row, 'Website', value, 'Added https:// for URL formatting only; destination not verified.')
        return candidate
    issue(file, sheet, row, 'Website', value, 'Invalid URL format: set NULL; original retained in audit.')
    return None

supplier_ids = {}
category_ids = {}
group_ids = {}
product_sources = defaultdict(list)
supplier_websites = defaultdict(list)

files = [
    ('Apromo - Retail Products-4.xlsx', 'Retail'),
    ('Apromo - Ingredients and Food Service Products-6.xlsx', 'Ingredients and Food Service'),
]
for filename, stream in files:
    for sheet, row, d in source_rows(filename):
        supplier = d['Supplier Name']
        if supplier not in supplier_ids:
            supplier_ids[supplier] = add('Supplier', len(supplier_ids) + 1, supplier, None, None)
        sid = supplier_ids[supplier]
        web = website(d.get('Website'), filename, sheet, row)
        if web and web not in supplier_websites[sid]:
            supplier_websites[sid].append(web)
            if DATA['Supplier'][sid - 1]['Website'] is None:
                DATA['Supplier'][sid - 1]['Website'] = web
            else:
                issue(filename, sheet, row, 'Website', web, 'Supplier has multiple URL strings; kept first well-formed URL as a provisional canonical value, not verified.')
        category = d.get('Category') or sheet
        if not d.get('Category'):
            issue(filename, sheet, row, 'Category', None, 'Used worksheet name as fallback category: ' + sheet)
        if category not in category_ids:
            category_ids[category] = add('ProductCategory', len(category_ids) + 1, category)
        groupkey = (stream, sheet)
        if groupkey not in group_ids:
            group_ids[groupkey] = add('ProductGroup', len(group_ids) + 1, sheet, stream)
        pid = len(DATA['Product']) + 1
        sourceid = d.get('Product ID')
        if not sourceid:
            issue(filename, sheet, row, 'Product ID', None, 'SourceProductID=NULL; generated internal ProductID=' + str(pid))
        else:
            product_sources[sourceid].append((filename, sheet, row, pid))
        price = Decimal(d.get('Unit Price') or d.get('Price'))
        assert price >= 0
        assert d['Discontinued'] in ('Yes', 'No')
        add('Product', pid, category_ids[category], group_ids[groupkey], sourceid,
            d.get('Product Description') or d['Product Descriptions'], 'Source pack',
            str(price.quantize(Decimal('.01'))), int(d['Discontinued'] == 'Yes'), None, None, None)
        add('ProductSupplier', pid, pid, sid, None, 1)
        LINEAGE.append([filename, sheet, row, 'Product', pid, 'One source row retained as one provisional product record.'])

for sourceid, refs in product_sources.items():
    if len(refs) > 1:
        for filename, sheet, row, pid in refs:
            issue(filename, sheet, row, 'Product ID', sourceid,
                  f'Non-unique source ID ({len(refs)} rows): kept separate provisional ProductID={pid}; identity/pricing requires review, no automatic merge.')

filename = 'Apromo - Distributors and Retailers-4.xlsx'
customer_ids = {}
site_ids = {}
region_by_sheet = {'NSW Retail stockists': 'NSW', 'QLD Stockists': 'QLD', 'VIC Retail Stockists': 'VIC', 'ACT Retail Stockists': 'ACT', 'NT Retail Stockists': 'NT'}
for sheet, row, d in source_rows(filename):
    name = d.get('Foodservice Distributor') or d.get('Retail stockist')
    assert name, (sheet, row)
    # Provisional organization grouping requires matching name AND normalized website host.
    web = website(d.get('Website'), filename, sheet, row)
    email = d.get('Email')
    if email and not re.fullmatch(r'[^\s@]+@[^\s@]+\.[^\s@]+', email):
        issue(filename, sheet, row, 'Email', email, 'Invalid email: set NULL; if clearly a URL and Website absent, move to Website.')
        if not web and '.' in email and ' ' not in email:
            web = website(email, filename, sheet, row)
        email = None
    host = urlparse(web).hostname.lower().removeprefix('www.') if web else None
    key = (name.casefold(), host) if host else (filename, sheet, row)
    ctype = 'Foodservice Distributor' if sheet == 'Foodservice Distributors' else 'Retail Stockist'
    if key not in customer_ids:
        customer_ids[key] = add('Customer', len(customer_ids) + 1, name, ctype, d.get('Description'))
    else:
        issue(filename, sheet, row, 'Customer', name, 'Provisional grouping by exact case-insensitive name plus website host; retained separate site; organization identity still needs review.')
    cid = customer_ids[key]
    region = region_by_sheet.get(sheet)
    country = 'New Zealand' if sheet == 'New Zealand Stockists' else 'Australia'
    address = d.get('Address')
    postcode = None
    # Only extract an explicit state + four-digit postcode; leave ambiguous values NULL.
    m = re.search(r'\b(NSW|VIC|QLD|SA|WA|TAS|ACT|NT)\s+(\d{4})\b', address or '', re.I)
    if m:
        if region and region != m[1].upper():
            issue(filename, sheet, row, 'StateRegion', address, 'Address conflicts with sheet region; retained sheet region pending review.')
        if not region and country == 'Australia':
            region = m[1].upper()
        if country == 'Australia':
            postcode = m[2]
    if sheet == 'New Zealand Stockists':
        issue(filename, sheet, row, 'Country/Address', address, 'New Zealand assigned from worksheet only; geography/phone values unverified and require review; no postcode inferred.')
    if d.get('Fax'):
        issue(filename, sheet, row, 'Fax', d['Fax'], 'No Fax field in approved schema; retained in raw source audit only.')
    siteid = len(DATA['CustomerSite']) + 1
    add('CustomerSite', siteid, cid, name, address, region, postcode, country, d.get('Phone'), email, web)
    site_ids[(sheet, row)] = siteid
    LINEAGE.append([filename, sheet, row, 'CustomerSite', siteid, 'Country/region inferred from sheet unless unambiguous address used; raw source retained.'])

# Fictional operational fixtures: four locations (two in each warehouse), five cases.
add('Warehouse', 1, 'Sydney - DEMO', 'Sydney', 'DEMO ADDRESS - Sydney, NSW, Australia')
add('Warehouse', 2, 'Melbourne - DEMO', 'Melbourne', 'DEMO ADDRESS - Melbourne, VIC, Australia')
for i, wid in [(1, 1), (2, 1), (3, 2), (4, 2), (5, 1)]:
    add('Worker', i, wid, f'DEMO Worker {i}', 'Full-time' if i <= 2 else 'Part-time')
for i, wid, code in [(1, 1, 'SYD-A01'), (2, 1, 'SYD-A02'), (3, 2, 'MEL-A01'), (4, 2, 'MEL-A02')]:
    add('StorageLocation', i, wid, code, 'DEMO chocolate zone', '15.00', '20.00', 1)

# Select five active, single-source-code products from the same supplier for clean fixture identity.
selected = [p for p in DATA['Product'] if p['Discontinued'] == 0 and p['SourceProductID'] and len(product_sources[p['SourceProductID']]) == 1 and DATA['ProductSupplier'][p['ProductID'] - 1]['SupplierID'] == supplier_ids['Butlers']][:5]
assert len(selected) == 5
demo_sites = [s['CustomerSiteID'] for s in DATA['CustomerSite'] if s['Country'] == 'Australia'][:5]
assert len(demo_sites) == 5
fixture_locations = [1, 1, 3, 3, 2]
received = [100, 100, 100, 100, 40]
shipped = [10, 18, 10, 0, 0]
picked = [10, 20, 10, 10, 5]
correct = [10, 18, 10, 10, None]
for i, product in enumerate(selected, 1):
    wid = 1 if i in (1, 2, 5) else 2
    worker = 1 if wid == 1 else 3
    location = fixture_locations[i - 1]
    received_at = '2026-03-04T09:00:00' if i != 2 else '2026-03-06T09:00:00'
    add('PurchaseOrder', i, supplier_ids['Butlers'], wid, '2026-03-01T09:00:00', '2026-03-05T17:00:00', 'Completed' if i < 5 else 'Partially Received')
    add('PurchaseOrderLine', i, i, product['ProductID'], '100.00', product['UnitPrice'])
    add('Receipt', i, i, worker, received_at, 'Completed')
    add('ReceiptLine', i, i, i, received[i - 1])
    add('InventoryLot', i, product['ProductID'], f'DEMO-LOT-{i:03}', '2027-03-31')
    add('ReceiptLineLot', i, i, i, received[i - 1])
    add('StockBalance', i, i, location, received[i - 1] - shipped[i - 1], '2026-03-12T12:00:00')
    order_status = 'Completed' if i in (1, 3) else ('Partially Delivered' if i == 2 else 'Processing')
    add('SalesOrder', i, demo_sites[i - 1], wid, '2026-03-07T09:00:00', '2026-03-10T17:00:00' if i < 5 else '2026-03-15T17:00:00', order_status)
    add('SalesOrderLine', i, i, product['ProductID'], picked[i - 1])
    add('PickLine', i, i, worker, i, picked[i - 1], correct[i - 1], '2026-03-08T10:00:00', '2026-03-08T10:30:00' if i < 5 else None)
    dispatched = '2026-03-09T09:00:00' if i <= 3 else None
    delivered = {1: '2026-03-10T12:00:00', 2: '2026-03-11T12:00:00', 3: '2026-03-10T14:00:00'}.get(i)
    add('Shipment', i, i, 'Own refrigerated van' if i != 3 else 'Cold-chain carrier', 'Kerry Logistics' if i == 3 else None, dispatched, delivered, '18.00' if delivered else None)
    if shipped[i - 1] > 0:
        add('ShipmentLine', i, i, i, shipped[i - 1])
        add('ShipmentLineLot', i, i, i, shipped[i - 1])

# Pending shipment headers are plans; no shipped line is recorded until quantity is confirmed.
for i, loc, temp in [(1, 1, '17.00'), (2, 2, '18.00'), (3, 3, '19.00'), (4, 4, '18.00'), (5, 1, '22.00')]:
    add('TemperatureLog', i, loc, '2026-03-12T10:00:00' if i < 5 else '2026-03-12T11:00:00', temp)

def write_csv(path, headers, rows):
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open('w', encoding='utf-8', newline='') as f:
        writer = csv.writer(f, lineterminator='\n')
        writer.writerow(headers)
        writer.writerows(rows)

for table, cols in SCHEMA.items():
    write_csv(OUT / 'csv' / (table + '.csv'), cols, ([r[c] for c in cols] for r in DATA[table]))
write_csv(OUT / 'audit' / 'cleaning_log.csv', ['SourceFile', 'Sheet', 'SourceRow', 'Field', 'OriginalValue', 'Action'], ISSUES)
write_csv(OUT / 'audit' / 'source_mapping.csv', ['SourceFile', 'Sheet', 'SourceRow', 'TargetTable', 'TargetID', 'Note'], LINEAGE)
(OUT / 'audit' / 'source_rows.json').write_text(json.dumps(RAW, ensure_ascii=False, indent=2), encoding='utf-8')
(OUT / 'schema_columns.json').write_text(json.dumps(SCHEMA, indent=2), encoding='utf-8')
summary = {
    'row_counts': {t: len(rows) for t, rows in DATA.items()},
    'product_source_rows': len(DATA['Product']),
    'duplicate_source_id_groups': sum(len(refs) > 1 for refs in product_sources.values()),
    'missing_product_source_ids': sum(p['SourceProductID'] is None for p in DATA['Product']),
    'cleaning_log_entries': len(ISSUES),
    'source_sha256': {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in SOURCE.glob('*.xlsx')},
    'fixture_product_ids': [p['ProductID'] for p in selected],
    'fixture_as_of': '2026-03-12T12:00:00',
    'expected_query_results': {
        'verified_pick_qty': 50, 'correct_pick_qty': 48, 'picking_accuracy_percent': '96.00',
        'delivered_shipments': 3, 'on_time_delivered_shipments': 2, 'delivered_on_time_percent': '66.67',
        'due_shipments': 4, 'overdue_undelivered_shipments': 1, 'due_shipment_on_time_percent': '50.00',
        'completed_purchase_orders': 4, 'on_time_completed_purchase_orders': 3, 'completed_receipts_on_time_percent': '75.00',
        'due_purchase_orders': 5, 'overdue_incomplete_purchase_orders': 1, 'due_order_on_time_receipt_percent': '60.00',
        'Sydney_occupied_locations': 2, 'Sydney_total_locations': 2, 'Sydney_utilization_percent': '100.00',
        'Melbourne_occupied_locations': 1, 'Melbourne_total_locations': 2, 'Melbourne_utilization_percent': '50.00',
    }
}
(OUT / 'audit' / 'manifest.json').write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding='utf-8')
print(json.dumps(summary, ensure_ascii=False, indent=2))
